from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Iterable, List

import pandas as pd
import soundfile as sf
from docx import Document as DocxDocument
from pypdf import PdfReader

from .models import Document


def ingest_text_file(path: Path, source: str = "text") -> List[Document]:
    text = path.read_text(encoding="utf-8")
    return [Document(id=path.name, text=text, metadata={"source": source, "path": str(path)})]


def ingest_jsonl(path: Path, source: str = "jsonl") -> List[Document]:
    docs: List[Document] = []
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            if not line.strip():
                continue
            raw = json.loads(line)
            docs.append(Document(**raw))
    return docs


def ingest_csv(path: Path, source: str = "csv") -> List[Document]:
    docs: List[Document] = []
    df = pd.read_csv(path)
    for idx, row in df.iterrows():
        text = json.dumps(row.to_dict(), ensure_ascii=False)
        docs.append(
            Document(id=f"{path.name}::{idx}", text=text, metadata={"source": source, "path": str(path)})
        )
    return docs


def ingest_excel(path: Path, source: str = "excel") -> List[Document]:
    docs: List[Document] = []
    xls = pd.ExcelFile(path)
    for sheet in xls.sheet_names:
        df = xls.parse(sheet)
        for idx, row in df.iterrows():
            text = json.dumps(row.to_dict(), ensure_ascii=False)
            docs.append(
                Document(
                    id=f"{path.name}::{sheet}::{idx}",
                    text=text,
                    metadata={"source": source, "path": str(path), "sheet": sheet},
                )
            )
    return docs


def ingest_pdf(path: Path, source: str = "pdf") -> List[Document]:
    reader = PdfReader(str(path))
    texts = []
    for i, page in enumerate(reader.pages):
        try:
            texts.append((i, page.extract_text() or ""))
        except Exception:
            texts.append((i, ""))
    docs = [Document(id=f"{path.name}::page{i}", text=txt, metadata={"source": source, "path": str(path)}) for i, txt in texts]
    return docs


def ingest_docx(path: Path, source: str = "docx") -> List[Document]:
    doc = DocxDocument(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    return [Document(id=f"{path.name}", text=text, metadata={"source": source, "path": str(path)})]


def ingest_audio(path: Path, source: str = "audio") -> List[Document]:
    # Placeholder: just capture metadata; real impl would call ASR
    info = sf.info(str(path))
    meta = {"source": source, "path": str(path), "samplerate": info.samplerate, "channels": info.channels}
    return [Document(id=f"{path.name}", text="", metadata=meta)]


def ingest_any(path: Path) -> List[Document]:
    suffix = path.suffix.lower()
    if suffix in {".txt"}:
        return ingest_text_file(path)
    if suffix in {".jsonl"}:
        return ingest_jsonl(path)
    if suffix in {".csv"}:
        return ingest_csv(path)
    if suffix in {".xls", ".xlsx"}:
        return ingest_excel(path)
    if suffix in {".pdf"}:
        return ingest_pdf(path)
    if suffix in {".docx"}:
        return ingest_docx(path)
    if suffix in {".wav", ".flac", ".ogg", ".mp3"}:
        return ingest_audio(path)
    raise ValueError(f"Unsupported file type: {suffix}")

